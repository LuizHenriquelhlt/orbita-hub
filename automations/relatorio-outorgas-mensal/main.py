"""Relatório mensal de outorgas de uso de água — automação de verdade.

Esta é a automação que roda **código real** na plataforma: lê o cadastro de
outorgas (CSV), consolida por região e situação, e produz duas saídas — a
planilha oficial em `.xlsx` e um CSV com as linhas rejeitadas.

Serve de referência para quem vai escrever a primeira automação, e por isso
exercita de propósito tudo que o contrato oferece: parâmetro validado, log em
nível, métrica de negócio, dois artefatos de formatos diferentes, e tratamento
de linha inválida sem abortar a execução.

## O dado

`dados/outorgas.csv` é sintético, gerado por `scripts/gerar_dados.py` a partir
de uma semente fixa. O domínio — Instituto Nacional de Recursos Hídricos — é
fictício. Nenhuma linha corresponde a outorga, processo ou pessoa real.

## Por que linha inválida não derruba a execução

Um ETL que aborta na primeira linha malformada é um ETL que nunca termina em
dados reais. Aqui, cada problema é contado, a linha vai para o CSV de rejeitos
com o motivo, e a execução termina com sucesso — mas com um `WARNING` visível
e a métrica `linhas_rejeitadas`, que é o que faz alguém olhar.

A exceção é a falha estrutural (arquivo ausente, coluna obrigatória faltando):
aí não há o que processar, e falhar é a resposta correta.
"""

from __future__ import annotations

import csv
import io
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from hub_sdk import Orbita

ARQUIVO_FONTE = Path(__file__).parent / "dados" / "outorgas.csv"

COLUNAS_OBRIGATORIAS = {
    "processo",
    "uf",
    "regiao",
    "finalidade",
    "situacao",
    "vazao_m3_s",
    "vencimento",
}

# Agrupamento por região, na ordem em que aparece no relatório impresso. Lista
# fixa, e não derivada do dado, porque uma região sem outorga no mês precisa
# aparecer com zero — e não desaparecer da planilha.
REGIOES = (
    "Norte",
    "Nordeste",
    "Centro-Oeste",
    "Sudeste",
    "Sul",
    "Bacia do São Francisco",
    "Bacia do Paraná",
    "Órgão Central",
)

SITUACOES = ("vigente", "em análise", "suspensa", "vencida")


@dataclass(slots=True)
class Outorga:
    processo: str
    uf: str
    regiao: str
    finalidade: str
    situacao: str
    vazao: Decimal
    vencimento: date


def main() -> None:
    with Orbita() as orbita:
        competencia = str(orbita.parametro("competencia") or _competencia_atual())
        regiao_filtro = orbita.parametro("regiao") or "todas"
        incluir_vencidas = bool(orbita.parametro("incluir_vencidas", False))

        orbita.info(f"competência {competencia}, região {regiao_filtro}")

        if not ARQUIVO_FONTE.is_file():
            # Falha estrutural: sem o arquivo não há o que processar.
            raise FileNotFoundError(
                f"cadastro de outorgas não encontrado em {ARQUIVO_FONTE}. "
                "Rode `python scripts/gerar_dados.py` para gerar o dado sintético."
            )

        orbita.info(f"lendo {ARQUIVO_FONTE.name}")
        validas, rejeitos = _carregar(orbita)

        orbita.metrica("linhas_lidas", len(validas) + len(rejeitos))
        orbita.metrica("linhas_validas", len(validas))
        orbita.metrica("linhas_rejeitadas", len(rejeitos))

        if rejeitos:
            orbita.aviso(f"{len(rejeitos)} linha(s) rejeitada(s) — ver outorgas-rejeitadas.csv")

        selecionadas = _filtrar(validas, regiao_filtro, incluir_vencidas)
        orbita.info(f"{len(selecionadas)} outorga(s) após os filtros")

        if not selecionadas:
            # Resultado vazio é resultado, não erro. Mas tem de ser dito com
            # clareza, senão alguém abre uma planilha em branco e acha que a
            # automação quebrou.
            orbita.aviso("nenhuma outorga atende aos filtros informados")
            orbita.resumo(f"Nenhuma outorga em {competencia} para {regiao_filtro}")
            return

        vazao_total = sum((o.vazao for o in selecionadas), Decimal(0))
        orbita.metrica("vazao_total_m3_s", float(vazao_total))
        orbita.metrica("regioes_com_dado", len({o.regiao for o in selecionadas}))

        planilha = _montar_planilha(selecionadas, competencia, orbita)
        orbita.artefato(f"outorgas-{competencia}.xlsx", planilha)

        if rejeitos:
            orbita.artefato("outorgas-rejeitadas.csv", _csv_de_rejeitos(rejeitos))

        vigentes = sum(1 for o in selecionadas if o.situacao == "vigente")
        orbita.resumo(
            f"{len(selecionadas)} outorgas ({vigentes} vigentes) em "
            f"{len({o.regiao for o in selecionadas})} regiões — "
            f"{vazao_total:.2f} m³/s"
        )


def _competencia_atual() -> str:
    hoje = date.today()
    return f"{hoje.year:04d}-{hoje.month:02d}"


def _carregar(orbita: Orbita) -> tuple[list[Outorga], list[dict[str, str]]]:
    """Lê o CSV, separando o que é aproveitável do que não é."""
    validas: list[Outorga] = []
    rejeitos: list[dict[str, str]] = []

    # `utf-8-sig` e não `utf-8`: o Excel grava CSV com BOM, e com `utf-8` o
    # BOM entra no nome da primeira coluna — que passa a ser `﻿processo`
    # e nunca casa com nada. É o bug mais comum ao ler planilha exportada.
    with ARQUIVO_FONTE.open(encoding="utf-8-sig", newline="") as arquivo:
        leitor = csv.DictReader(arquivo, delimiter=";")

        faltando = COLUNAS_OBRIGATORIAS - set(leitor.fieldnames or [])
        if faltando:
            raise ValueError(f"colunas obrigatórias ausentes no cadastro: {sorted(faltando)}")

        for numero, linha in enumerate(leitor, start=2):
            try:
                validas.append(_interpretar(linha))
            except (ValueError, InvalidOperation, KeyError) as erro:
                rejeitos.append({**linha, "_linha": str(numero), "_motivo": str(erro)})
                if len(rejeitos) <= 5:
                    # Só as cinco primeiras no log: um arquivo com 2000 linhas
                    # ruins não deve produzir 2000 linhas de log.
                    orbita.debug(f"linha {numero} rejeitada: {erro}")

    return validas, rejeitos


def _interpretar(linha: dict[str, str]) -> Outorga:
    vazao_bruta = (linha.get("vazao_m3_s") or "").strip().replace(",", ".")
    if not vazao_bruta:
        raise ValueError("vazão vazia")
    vazao = Decimal(vazao_bruta)
    if vazao < 0:
        raise ValueError(f"vazão negativa: {vazao}")

    situacao = (linha.get("situacao") or "").strip().lower()
    if situacao not in SITUACOES:
        raise ValueError(f"situação desconhecida: {situacao!r}")

    vencimento_bruto = (linha.get("vencimento") or "").strip()
    try:
        vencimento = date.fromisoformat(vencimento_bruto)
    except ValueError as erro:
        raise ValueError(f"vencimento inválido: {vencimento_bruto!r}") from erro

    processo = (linha.get("processo") or "").strip()
    if not processo:
        raise ValueError("processo vazio")

    return Outorga(
        processo=processo,
        uf=(linha.get("uf") or "").strip().upper(),
        regiao=(linha.get("regiao") or "").strip(),
        finalidade=(linha.get("finalidade") or "").strip(),
        situacao=situacao,
        vazao=vazao,
        vencimento=vencimento,
    )


def _filtrar(outorgas: list[Outorga], regiao: str, incluir_vencidas: bool) -> list[Outorga]:
    selecionadas = outorgas
    if regiao and regiao != "todas":
        selecionadas = [o for o in selecionadas if o.regiao == regiao]
    if not incluir_vencidas:
        selecionadas = [o for o in selecionadas if o.situacao != "vencida"]
    return selecionadas


def _csv_de_rejeitos(rejeitos: list[dict[str, str]]) -> bytes:
    """CSV das linhas rejeitadas, com o motivo na primeira coluna.

    Motivo primeiro porque é a informação que quem recebe o arquivo precisa —
    e porque, em uma planilha larga, a última coluna é a que ninguém rola para
    ver.
    """
    campos = ["_motivo", "_linha", *sorted(set(rejeitos[0]) - {"_motivo", "_linha"})]
    buffer = io.StringIO(newline="")
    escritor = csv.DictWriter(
        buffer, fieldnames=campos, delimiter=";", lineterminator="\r\n", extrasaction="ignore"
    )
    escritor.writeheader()
    escritor.writerows(rejeitos)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def _montar_planilha(outorgas: list[Outorga], competencia: str, orbita: Orbita) -> bytes:
    """Gera o `.xlsx` com três abas: resumo, por região e detalhe."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    planilha = Workbook()

    azul = PatternFill("solid", fgColor="1F3A5F")
    branco_negrito = Font(color="FFFFFF", bold=True)
    borda_fina = Border(bottom=Side(style="thin", color="D0D7DE"))

    def cabecalho(aba: object, titulos: list[str], larguras: list[int]) -> None:
        for indice, (titulo, largura) in enumerate(zip(titulos, larguras, strict=True), start=1):
            celula = aba.cell(row=1, column=indice, value=titulo)  # type: ignore[attr-defined]
            celula.fill = azul
            celula.font = branco_negrito
            celula.alignment = Alignment(horizontal="center", vertical="center")
            aba.column_dimensions[get_column_letter(indice)].width = largura  # type: ignore[attr-defined]
        aba.freeze_panes = "A2"  # type: ignore[attr-defined]

    # ── Aba 1: resumo por situação ───────────────────────────────────────────
    resumo = planilha.active
    resumo.title = "Resumo"
    cabecalho(resumo, ["Situação", "Outorgas", "Vazão (m³/s)", "% do total"], [22, 12, 16, 12])

    por_situacao: Counter[str] = Counter()
    vazao_por_situacao: defaultdict[str, Decimal] = defaultdict(lambda: Decimal(0))
    for outorga in outorgas:
        por_situacao[outorga.situacao] += 1
        vazao_por_situacao[outorga.situacao] += outorga.vazao

    total = len(outorgas)
    for indice, situacao in enumerate(SITUACOES, start=2):
        quantidade = por_situacao.get(situacao, 0)
        resumo.cell(row=indice, column=1, value=situacao.capitalize())
        resumo.cell(row=indice, column=2, value=quantidade)
        resumo.cell(
            row=indice, column=3, value=float(vazao_por_situacao[situacao])
        ).number_format = "#,##0.000"
        resumo.cell(
            row=indice, column=4, value=(quantidade / total if total else 0)
        ).number_format = "0.0%"

    linha_total = len(SITUACOES) + 2
    resumo.cell(row=linha_total, column=1, value="Total").font = Font(bold=True)
    # Fórmula, não valor calculado: quem abre a planilha vê de onde o número
    # saiu e pode conferir. Relatório oficial com total "mágico" gera chamado.
    resumo.cell(row=linha_total, column=2, value=f"=SUM(B2:B{linha_total - 1})").font = Font(
        bold=True
    )
    resumo.cell(row=linha_total, column=3, value=f"=SUM(C2:C{linha_total - 1})").font = Font(
        bold=True
    )
    resumo.cell(row=linha_total, column=3).number_format = "#,##0.000"

    resumo["F1"] = "Competência"
    resumo["G1"] = competencia
    resumo["F2"] = "Gerado por"
    resumo["G2"] = f"Órbita — execução {orbita.execucao_id[:8]}"
    resumo["F3"] = "Origem dos dados"
    resumo["G3"] = "Sintético (INRH é um órgão fictício)"
    for celula in ("F1", "F2", "F3"):
        resumo[celula].font = Font(bold=True)
    resumo.column_dimensions["F"].width = 18
    resumo.column_dimensions["G"].width = 42

    # ── Aba 2: por região ────────────────────────────────────────────────────
    por_regiao = planilha.create_sheet("Por região")
    cabecalho(
        por_regiao,
        ["Região", "Outorgas", "Vigentes", "Vazão (m³/s)", "Vazão média"],
        [26, 12, 12, 16, 14],
    )

    agrupado: defaultdict[str, list[Outorga]] = defaultdict(list)
    for outorga in outorgas:
        agrupado[outorga.regiao].append(outorga)

    for indice, regiao in enumerate(REGIOES, start=2):
        itens = agrupado.get(regiao, [])
        vazao = sum((o.vazao for o in itens), Decimal(0))
        por_regiao.cell(row=indice, column=1, value=regiao)
        por_regiao.cell(row=indice, column=2, value=len(itens))
        por_regiao.cell(
            row=indice, column=3, value=sum(1 for o in itens if o.situacao == "vigente")
        )
        por_regiao.cell(row=indice, column=4, value=float(vazao)).number_format = "#,##0.000"
        por_regiao.cell(
            row=indice, column=5, value=float(vazao / len(itens)) if itens else 0
        ).number_format = "#,##0.000"

    # ── Aba 3: detalhe ───────────────────────────────────────────────────────
    detalhe = planilha.create_sheet("Detalhe")
    cabecalho(
        detalhe,
        ["Processo", "UF", "Região", "Finalidade", "Situação", "Vazão (m³/s)", "Vencimento"],
        [24, 6, 26, 26, 14, 14, 14],
    )

    for indice, outorga in enumerate(sorted(outorgas, key=lambda o: (o.regiao, -o.vazao)), start=2):
        detalhe.cell(row=indice, column=1, value=outorga.processo).border = borda_fina
        detalhe.cell(row=indice, column=2, value=outorga.uf)
        detalhe.cell(row=indice, column=3, value=outorga.regiao)
        detalhe.cell(row=indice, column=4, value=outorga.finalidade)
        detalhe.cell(row=indice, column=5, value=outorga.situacao)
        detalhe.cell(row=indice, column=6, value=float(outorga.vazao)).number_format = "#,##0.000"
        detalhe.cell(row=indice, column=7, value=outorga.vencimento).number_format = "dd/mm/yyyy"

    detalhe.auto_filter.ref = f"A1:G{len(outorgas) + 1}"

    propriedades = planilha.properties
    propriedades.title = f"Outorgas {competencia}"
    propriedades.creator = "Órbita — dado sintético"

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    main()
