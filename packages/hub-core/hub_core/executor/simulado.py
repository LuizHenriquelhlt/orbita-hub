"""Executor simulado: reproduz um roteiro em vez de rodar código de verdade.

## Por que isto existe, e por que não é trapaça

A demonstração pública da plataforma fica em um container de 512 MB sem acesso
ao Docker Engine. Rodar container arbitrário ali é impossível — e, se fosse
possível, seria imprudente: um endereço público que executa imagem informada
pelo visitante é um serviço de mineração de criptomoeda à espera de ser
descoberto.

O que o visitante precisa ver é o *mecanismo da plataforma*: a execução entrar
na fila, sair dela, emitir log em tempo real, publicar métrica, produzir um
arquivo que se baixa, falhar e ser retentada com backoff. Tudo isso é código de
verdade — fila, motor, callbacks, armazenamento. O único trecho encenado é o
que normalmente estaria dentro do container de terceiro.

A alternativa honesta é deixar claro qual executor rodou cada execução, e é o
que a plataforma faz: `execucao.executor` é persistido e aparece na interface,
e o roteiro de cada automação simulada está no manifesto, em texto, no
repositório. Nada é apresentado como real sem ser.

## O roteiro

Uma lista de passos no manifesto da automação. Cada passo pode esperar, emitir
log, publicar métrica, gerar artefato ou falhar:

    roteiro:
      - espera_ms: 400
        log: "Conectando ao cadastro de outorgas"
      - espera_ms: 1200
        log: "1420 registros lidos"
        metrica: { linhas_lidas: 1420 }
      - artefato:
          caminho: outorgas-2026-09.csv
          formato: csv
          colunas: [uf, processo, vazao_m3_s, situacao]
          linhas: 60
      - resumo: "1420 outorgas consolidadas"

Os dados do artefato são gerados por um PRNG semeado pelo id da execução, então
duas execuções da mesma automação produzem arquivos diferentes — como
aconteceria de verdade — mas uma execução reexecutada com o mesmo id produz o
mesmo arquivo, o que torna o teste determinístico.
"""

from __future__ import annotations

import asyncio
import csv
import io
import random
from datetime import UTC, datetime, timedelta
from typing import Any

from hub_core.executor.base import (
    AoRegistrarLog,
    ArtefatoProduzido,
    Executor,
    PedidoDeExecucao,
    ResultadoExecucao,
    StatusExecucao,
)
from hub_core.logging import obter_logger

log = obter_logger(__name__)

# Vocabulário do domínio fictício, usado para preencher os artefatos gerados.
_UFS = ("AM", "BA", "CE", "GO", "MG", "MT", "PA", "PE", "PR", "RS", "SP", "TO")
_SITUACOES = ("vigente", "em análise", "suspensa", "vencida")
_FINALIDADES = (
    "abastecimento urbano",
    "irrigação",
    "indústria",
    "mineração",
    "aquicultura",
    "termoelétrica",
)


class ExecutorSimulado(Executor):
    nome = "simulado"

    def __init__(self, *, acelerar: bool = False) -> None:
        self._acelerar = acelerar
        """Em teste, `acelerar=True` zera as esperas e a suíte roda em milissegundos."""
        self._cancelados: set[str] = set()

    async def executar(
        self,
        pedido: PedidoDeExecucao,
        ao_registrar_log: AoRegistrarLog,
    ) -> ResultadoExecucao:
        aleatorio = random.Random(pedido.execucao_id)
        metricas: dict[str, float] = {}
        artefatos: list[ArtefatoProduzido] = []
        resumo = ""
        inicio = datetime.now(UTC)

        roteiro = pedido.roteiro_simulado or _roteiro_padrao()

        await ao_registrar_log(
            "INFO",
            f"[simulado] iniciando {pedido.automacao_slug} "
            f"({len(roteiro)} passos) — executor sem container",
        )

        for indice, passo in enumerate(roteiro, start=1):
            if pedido.execucao_id in self._cancelados:
                self._cancelados.discard(pedido.execucao_id)
                await ao_registrar_log("WARNING", "[simulado] cancelada a pedido do usuário")
                return ResultadoExecucao(
                    status=StatusExecucao.CANCELADA,
                    resumo="Cancelada durante a execução",
                    metricas=metricas,
                    artefatos=artefatos,
                    consumo=self._consumo(inicio, aleatorio),
                )

            espera_ms = int(passo.get("espera_ms", 300))
            if not self._acelerar and espera_ms > 0:
                await asyncio.sleep(espera_ms / 1000)

            if mensagem := passo.get("log"):
                await ao_registrar_log(str(passo.get("nivel", "INFO")), str(mensagem))

            for nome, valor in (passo.get("metrica") or {}).items():
                metricas[str(nome)] = float(valor)
                await ao_registrar_log("DEBUG", f"métrica {nome}={valor}")

            if especificacao := passo.get("artefato"):
                artefato = _gerar_artefato(especificacao, aleatorio)
                artefatos.append(artefato)
                await ao_registrar_log(
                    "INFO",
                    f"artefato produzido: {artefato.caminho_relativo} "
                    f"({len(artefato.conteudo)} bytes)",
                )

            if texto := passo.get("resumo"):
                resumo = str(texto)

            if falha := passo.get("falha"):
                await ao_registrar_log("ERROR", str(falha))
                return ResultadoExecucao(
                    status=StatusExecucao.FALHA,
                    codigo_saida=int(passo.get("codigo_saida", 1)),
                    resumo=resumo or "Falhou durante a execução",
                    erro=str(falha),
                    metricas=metricas,
                    artefatos=artefatos,
                    consumo=self._consumo(inicio, aleatorio),
                )

            if passo.get("timeout"):
                await ao_registrar_log("ERROR", "tempo máximo de execução excedido")
                return ResultadoExecucao(
                    status=StatusExecucao.TIMEOUT,
                    resumo="Tempo esgotado",
                    erro=f"excedeu {pedido.limites.timeout_segundos}s",
                    metricas=metricas,
                    artefatos=artefatos,
                    consumo=self._consumo(inicio, aleatorio),
                )

            log.debug("simulado.passo", execucao=pedido.execucao_id, passo=indice)

        return ResultadoExecucao(
            status=StatusExecucao.SUCESSO,
            codigo_saida=0,
            resumo=resumo or "Execução concluída",
            metricas=metricas,
            artefatos=artefatos,
            consumo=self._consumo(inicio, aleatorio),
        )

    async def cancelar(self, execucao_id: str) -> bool:
        self._cancelados.add(execucao_id)
        return True

    async def saudavel(self) -> tuple[bool, str]:
        return True, "executor simulado — não requer Docker"

    def _consumo(self, inicio: datetime, aleatorio: random.Random) -> dict[str, Any]:
        """Números de consumo plausíveis, marcados como simulados.

        A chave `simulado: True` viaja com o dado até a interface, que a usa
        para rotular o número. Métrica inventada sem rótulo é a única parte
        realmente desonesta de uma demonstração.
        """
        duracao = (datetime.now(UTC) - inicio).total_seconds()
        return {
            "simulado": True,
            "duracao_segundos": round(duracao, 2),
            "cpu_segundos": round(duracao * aleatorio.uniform(0.15, 0.6), 2),
            "memoria_pico_mb": aleatorio.randint(60, 240),
        }


def _roteiro_padrao() -> list[dict[str, Any]]:
    """Usado quando a automação não declara roteiro: um ciclo curto de sucesso."""
    return [
        {"espera_ms": 400, "log": "Preparando ambiente de execução"},
        {"espera_ms": 900, "log": "Processando", "metrica": {"itens": 1}},
        {"espera_ms": 300, "resumo": "Execução concluída"},
    ]


def _gerar_artefato(especificacao: dict[str, Any], aleatorio: random.Random) -> ArtefatoProduzido:
    caminho = str(especificacao.get("caminho", "saida.txt"))
    formato = str(especificacao.get("formato", "texto")).lower()
    linhas = int(especificacao.get("linhas", 25))
    colunas = [str(c) for c in especificacao.get("colunas", ["chave", "valor"])]

    if formato == "csv":
        return ArtefatoProduzido(
            caminho_relativo=caminho,
            conteudo=_csv_sintetico(colunas, linhas, aleatorio),
            tipo_conteudo="text/csv; charset=utf-8",
        )
    if formato == "xlsx":
        return ArtefatoProduzido(
            caminho_relativo=caminho,
            conteudo=_xlsx_sintetico(colunas, linhas, aleatorio, titulo=caminho),
            tipo_conteudo=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        )

    texto = especificacao.get("conteudo") or "\n".join(
        f"linha {n + 1}: {aleatorio.randint(1000, 9999)}" for n in range(linhas)
    )
    return ArtefatoProduzido(
        caminho_relativo=caminho,
        conteudo=str(texto).encode("utf-8"),
        tipo_conteudo="text/plain; charset=utf-8",
    )


def _valor_para(coluna: str, indice: int, aleatorio: random.Random) -> Any:
    """Gera um valor plausível a partir do nome da coluna.

    Mapear nome para tipo de dado é o que separa um artefato que parece real de
    um cheio de `lorem ipsum`. Quem abre a planilha da demonstração vê UF em
    coluna de UF e vazão em m³/s.
    """
    nome = coluna.lower()
    if "uf" in nome or "estado" in nome:
        return aleatorio.choice(_UFS)
    if "processo" in nome or "protocolo" in nome:
        return f"{aleatorio.randint(10000, 99999)}.{aleatorio.randint(2019, 2026)}-{indice:04d}"
    if "situacao" in nome or "status" in nome:
        return aleatorio.choice(_SITUACOES)
    if "finalidade" in nome or "uso" in nome:
        return aleatorio.choice(_FINALIDADES)
    if "vazao" in nome or "volume" in nome:
        return round(aleatorio.uniform(0.01, 48.5), 3)
    if "valor" in nome or "custo" in nome:
        return round(aleatorio.uniform(1200, 980000), 2)
    if "data" in nome or "vencimento" in nome:
        base = datetime(2026, 1, 1, tzinfo=UTC) + timedelta(days=aleatorio.randint(0, 900))
        return base.date().isoformat()
    if "quantidade" in nome or "total" in nome:
        return aleatorio.randint(1, 5000)
    return f"{coluna}-{indice:04d}"


def _csv_sintetico(colunas: list[str], linhas: int, aleatorio: random.Random) -> bytes:
    buffer = io.StringIO(newline="")
    escritor = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    escritor.writerow(colunas)
    for indice in range(linhas):
        escritor.writerow([_valor_para(coluna, indice, aleatorio) for coluna in colunas])
    # BOM para que o Excel em português abra o arquivo com acentos corretos sem
    # o usuário ter de escolher a codificação na importação.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def _xlsx_sintetico(
    colunas: list[str], linhas: int, aleatorio: random.Random, *, titulo: str
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Dados"

    cabecalho_fundo = PatternFill("solid", fgColor="1F3A5F")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)

    for coluna_idx, coluna in enumerate(colunas, start=1):
        celula = aba.cell(row=1, column=coluna_idx, value=coluna.replace("_", " ").title())
        celula.fill = cabecalho_fundo
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center", vertical="center")
        aba.column_dimensions[get_column_letter(coluna_idx)].width = max(14, len(coluna) + 6)

    for linha_idx in range(linhas):
        for coluna_idx, coluna in enumerate(colunas, start=1):
            aba.cell(
                row=linha_idx + 2,
                column=coluna_idx,
                value=_valor_para(coluna, linha_idx, aleatorio),
            )

    # Congelar a primeira linha e ligar o autofiltro: é o que distingue uma
    # planilha gerada por programa de uma planilha utilizável.
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{linhas + 1}"

    propriedades = planilha.properties
    propriedades.title = titulo
    propriedades.creator = "Órbita (dado sintético)"

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()
